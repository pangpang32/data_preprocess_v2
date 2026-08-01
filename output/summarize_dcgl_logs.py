"""
summarize_dcgl_logs.py

汇总 DCGL 对比实验日志（MINE/logs/full_logs/{dataset}/*.log）：

每条日志解析两个信息：
    第 2 行：DESC: {desc}（如 mine_kg_4660_l2_3）
    最后一个非空行：early stopping at {N}, recall@50:{R}（最佳轮 = N - 10，patience=10）

输出（写入日志根目录）：
    汇总.txt   # 分组明细 + 按 mean recall 降序的总表
    汇总.xlsx  # 汇总 / 明细 / 透视 三个 sheet

执行位置：项目根目录（data_preprocess/）
用法：python output/summarize_dcgl_logs.py [--log_root ../MINE/logs/full_logs]
"""
import os
import re
import argparse
from collections import defaultdict

PATIENCE = 10  # early stopping at N -> 最佳轮 = N - PATIENCE

LINE_PAT = re.compile(r"early stopping at (\d+),\s*recall@50:([\d.]+)")
DESC_PAT = re.compile(r"^DESC:\s*(\S+)")


def parse_dataset(folder: str):
    """文件夹名 -> (group, kg_version, norm)。"""
    parts = folder.split("_")
    norm = "no_l2" if folder.endswith("no_l2") else "l2"
    group = "物品嵌入" if folder.startswith("item_") else "全量嵌入"
    if "catfreq" in folder:
        kg = "catfreq" + folder.split("catfreq")[1].split("_")[0]
    elif "4660" in folder:
        kg = "4660"
    else:
        kg = "原始KG"
    return group, kg, norm


def parse_log(path: str):
    desc = None
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        lines = f.readlines()
    for line in lines[:5]:
        m = DESC_PAT.match(line.strip())
        if m:
            desc = m.group(1)
            break
    for line in reversed(lines):
        line = line.strip()
        if not line:
            continue
        m = LINE_PAT.search(line)
        if m:
            return desc, int(m.group(1)) - PATIENCE, float(m.group(2))
        return desc, None, None
    return desc, None, None


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--log_root", type=str, default="../MINE/logs/full_logs")
    args = parser.parse_args()
    root = args.log_root

    # records: (dataset, group, kg, norm, run, best_epoch, recall)
    records = []
    for folder in sorted(os.listdir(root)):
        fdir = os.path.join(root, folder)
        if not os.path.isdir(fdir):
            continue
        group, kg, norm = parse_dataset(folder)
        for fname in sorted(os.listdir(fdir)):
            if not fname.endswith(".log"):
                continue
            desc, best_epoch, recall = parse_log(os.path.join(fdir, fname))
            run = None
            if desc:
                m = re.search(r"_(\d+)$", desc)
                if m:
                    run = int(m.group(1))
            records.append({
                "dataset": folder, "group": group, "kg": kg, "norm": norm,
                "run": run, "best_epoch": best_epoch, "recall": recall,
                "file": fname,
            })

    # 按 dataset+norm 聚合
    agg = defaultdict(list)
    for r in records:
        agg[(r["dataset"], r["group"], r["kg"], r["norm"])].append(r)
    for key in agg:
        agg[key].sort(key=lambda r: (r["run"] is None, r["run"]))

    def mean(xs):
        return sum(xs) / len(xs) if xs else float("nan")

    def std(xs):
        if len(xs) < 2:
            return 0.0
        m = mean(xs)
        return (sum((x - m) ** 2 for x in xs) / (len(xs) - 1)) ** 0.5

    summaries = []
    for (dataset, group, kg, norm), runs in agg.items():
        recalls = [r["recall"] for r in runs if r["recall"] is not None]
        epochs = [r["best_epoch"] for r in runs if r["best_epoch"] is not None]
        summaries.append({
            "dataset": dataset, "group": group, "kg": kg, "norm": norm,
            "runs": runs,
            "recalls": recalls, "epochs": epochs,
            "mean": mean(recalls), "std": std(recalls),
            "mean_epoch": mean(epochs),
        })

    kg_order = {"原始KG": 0, "4660": 1, "catfreq2": 2, "catfreq3": 3, "catfreq4": 4, "catfreq5": 5}
    group_order = {"物品嵌入": 0, "全量嵌入": 1}
    summaries.sort(key=lambda s: (kg_order.get(s["kg"], 9), group_order.get(s["group"], 9), s["norm"]))

    # ---------------- 汇总.txt ----------------
    txt_lines = ["DCGL 对比实验汇总（recall@50；3 次运行 mean±std；轮次 = 最佳轮）", ""]
    cur_group = None
    for s in summaries:
        head = f"{s['kg']} + {s['group']}"
        if head != cur_group:
            txt_lines.append(head)
            cur_group = head
        run_str = "  ".join(f"{r['recall']:.4f}" if r["recall"] is not None else "FAILED" for r in s["runs"])
        ep_str = ",".join(str(r["best_epoch"]) if r["best_epoch"] is not None else "-" for r in s["runs"])
        txt_lines.append(
            f" - {s['norm']:<6} {s['mean']:.4f} ± {s['std']:.4f}  (runs: {run_str}; 轮次: {ep_str})"
        )
    txt_lines += ["", "按 mean recall@50 降序：", ""]
    for i, s in enumerate(sorted(summaries, key=lambda x: -x["mean"]), 1):
        txt_lines.append(f"{i:>2}. {s['dataset']:<32} {s['mean']:.4f} ± {s['std']:.4f}")

    txt_path = os.path.join(root, "汇总.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("\n".join(txt_lines) + "\n")
    print(f"Wrote {txt_path}")

    # ---------------- 汇总.xlsx ----------------
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")

    def style_sheet(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
        ws.freeze_panes = "A2"

    # Sheet 1: 汇总
    ws = wb.active
    ws.title = "汇总"
    ws.append(["数据集", "分组", "KG版本", "归一化",
               "run1", "run2", "run3", "mean", "std",
               "轮次1", "轮次2", "轮次3", "平均最佳轮"])
    for s in summaries:
        runs = s["runs"] + [None] * (3 - len(s["runs"]))
        row = [s["dataset"], s["group"], s["kg"], s["norm"]]
        row += [r["recall"] if r and r["recall"] is not None else None for r in runs]
        row += [s["mean"], s["std"]]
        row += [r["best_epoch"] if r and r["best_epoch"] is not None else None for r in runs]
        row += [round(s["mean_epoch"], 1) if s["epochs"] else None]
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(5, 10):
            ws.cell(row=r, column=c).number_format = "0.0000"
    style_sheet(ws, 13)
    for i, w in enumerate([30, 10, 10, 8, 9, 9, 9, 9, 8, 7, 7, 7, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Sheet 2: 明细（tidy，便于透视/画图）
    ws2 = wb.create_sheet("明细")
    ws2.append(["数据集", "分组", "KG版本", "归一化", "run", "recall@50", "最佳轮", "日志文件"])
    for r in records:
        ws2.append([r["dataset"], r["group"], r["kg"], r["norm"], r["run"],
                    r["recall"], r["best_epoch"], r["file"]])
    for r in range(2, ws2.max_row + 1):
        ws2.cell(row=r, column=6).number_format = "0.0000"
    style_sheet(ws2, 8)
    for i, w in enumerate([30, 10, 10, 8, 6, 10, 8, 45], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # Sheet 3: 透视（行=KG版本，列=分组×归一化，值=mean recall@50）
    ws3 = wb.create_sheet("透视")
    norms = ["l2", "no_l2"]
    groups = ["物品嵌入", "全量嵌入"]
    ws3.append(["KG版本"] + [f"{g}-{n}" for g in groups for n in norms])
    lookup = {(s["kg"], s["group"], s["norm"]): s["mean"] for s in summaries}
    for kg in ["原始KG", "4660", "catfreq2", "catfreq3", "catfreq4", "catfreq5"]:
        row = [kg]
        for g in groups:
            for n in norms:
                row.append(lookup.get((kg, g, n)))
        ws3.append(row)
    for r in range(2, ws3.max_row + 1):
        for c in range(2, 6):
            ws3.cell(row=r, column=c).number_format = "0.0000"
    style_sheet(ws3, 5)
    for i, w in enumerate([10, 16, 16, 16, 16], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    xlsx_path = os.path.join(root, "汇总.xlsx")
    wb.save(xlsx_path)
    print(f"Wrote {xlsx_path}")
    print(f"共 {len(records)} 条日志，{len(summaries)} 组")


if __name__ == "__main__":
    main()
